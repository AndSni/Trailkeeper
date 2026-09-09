"""Web console (server-rendered) - auth pages + dashboard."""


def _register(client, email="console@example.com"):
    return client.post(
        "/register",
        data={
            "organisation_name": "Console Crew",
            "name": "Cass",
            "email": email,
            "password": "a decent long passphrase",
        },
        follow_redirects=False,
    )


def test_index_and_login_page(client):
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"] == "/login"

    page = client.get("/login")
    assert page.status_code == 200
    assert "Sign in" in page.text


def test_app_requires_session(client):
    r = client.get("/app", follow_redirects=False)
    assert r.status_code in (303, 307)
    assert r.headers["location"] == "/login"


def test_register_sets_session_and_dashboard_loads(client):
    r = _register(client)
    assert r.status_code == 303
    assert r.headers["location"] == "/app"
    assert "tk_session" in r.cookies

    dash = client.get("/app")  # client keeps the cookie
    assert dash.status_code == 200
    assert "Console Crew" in dash.text
    assert "Structures by status" in dash.text


def test_dashboard_shows_project_sections(client):
    _register(client, email="p@example.com")
    # the console's session cookie also authenticates the JSON API? no - use the
    # API with its own token. Grab one via the JSON login.
    tokens = client.post(
        "/auth/login",
        json={"email": "p@example.com", "password": "a decent long passphrase"},
    ).json()
    h = {"Authorization": f"Bearer {tokens['access_token']}"}
    client.post("/projects", headers=h, json={"name": "Blue Trail", "activity": "mtb"})

    dash = client.get("/app")
    assert dash.status_code == 200
    assert "Blue Trail" in dash.text
    assert "Job-type productivity" in dash.text
    assert "Hours per member" in dash.text


def test_login_bad_credentials_rerenders(client):
    _register(client, email="real@example.com")
    client.get("/logout")
    r = client.post(
        "/login",
        data={"email": "real@example.com", "password": "wrong"},
        follow_redirects=False,
    )
    assert r.status_code == 200
    assert "Invalid email or password" in r.text


def test_members_page_and_invite_creation(client):
    _register(client, email="admin@example.com")

    members = client.get("/app/members")
    assert members.status_code == 200
    assert "admin@example.com" in members.text

    made = client.post(
        "/app/invites",
        data={"email": "newbie@example.com", "org_role": "editor"},
        follow_redirects=False,
    )
    assert made.status_code == 303
    assert "/app/members?invited=" in made.headers["location"]
    token = made.headers["location"].split("invited=")[1]

    invite_page = client.get(f"/invite/{token}")
    assert invite_page.status_code == 200
    assert "Join Console Crew" in invite_page.text

    accepted = client.post(
        f"/invite/{token}",
        data={"name": "Newbie", "password": "another good passphrase"},
        follow_redirects=False,
    )
    assert accepted.status_code == 303
    assert accepted.headers["location"] == "/app"
    assert "tk_session" in accepted.cookies


def test_invite_page_invalid_token(client):
    page = client.get("/invite/not-a-real-token")
    assert page.status_code == 200
    assert "Invite not valid" in page.text


def test_exports_require_session(client):
    r = client.get("/app/export/hours.csv", follow_redirects=False)
    assert r.status_code in (303, 307)
    x = client.get("/app/export.xlsx", follow_redirects=False)
    assert x.status_code in (303, 307)


def test_csv_and_xlsx_exports(client):
    _register(client, email="ex@example.com")
    tokens = client.post(
        "/auth/login",
        json={"email": "ex@example.com", "password": "a decent long passphrase"},
    ).json()
    h = {"Authorization": f"Bearer {tokens['access_token']}"}
    pid = client.post(
        "/projects", headers=h, json={"name": "Exportable", "activity": "mtb"}
    ).json()["id"]
    client.post(
        f"/tasks?project_id={pid}",
        headers=h,
        json={"title": "Fix the berm", "priority": "high", "estimate_min": 45},
    )
    struct_id = client.post(
        "/structures", headers=h, json={"name": "Culvert 9", "structure_type": "culvert"}
    ).json()["id"]
    client.post(
        "/inspections",
        headers=h,
        json={"project_id": pid, "structure_id": struct_id, "risk": "medium"},
    )

    tasks_csv = client.get("/app/export/tasks.csv")
    assert tasks_csv.status_code == 200
    assert tasks_csv.headers["content-type"].startswith("text/csv")
    assert "attachment" in tasks_csv.headers["content-disposition"]
    body = tasks_csv.text
    assert body.splitlines()[0].startswith("Title,Type,Priority,Status")
    assert "Fix the berm" in body

    ins_csv = client.get("/app/export/inspections.csv")
    assert ins_csv.status_code == 200
    assert "Culvert 9" in ins_csv.text
    assert "medium" in ins_csv.text

    bad = client.get("/app/export/nonsense.csv")
    assert bad.status_code == 404

    xlsx = client.get("/app/export.xlsx")
    assert xlsx.status_code == 200
    assert "spreadsheetml" in xlsx.headers["content-type"]
    assert xlsx.content[:2] == b"PK"  # xlsx is a zip

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xlsx.content))
    assert set(wb.sheetnames) == {
        "Hours", "Productivity", "Tasks", "Segments", "Structures", "Inspections"
    }
    assert wb["Tasks"].cell(row=1, column=1).value == "Title"
    task_titles = [row[0] for row in wb["Tasks"].iter_rows(min_row=2, values_only=True)]
    assert "Fix the berm" in task_titles
